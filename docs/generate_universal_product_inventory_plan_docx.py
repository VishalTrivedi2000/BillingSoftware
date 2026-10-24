"""
Generate Universal Product & Inventory Management — Implementation Plan (.docx).

Targets MS SQL Server (same as Billing workspace). Describes extending the existing
product/catalog implementation toward a metadata-driven universal model + ledger inventory.

Output: docs/Universal_Product_Inventory_Management_Plan.docx

Run from repo root:
  python docs/generate_universal_product_inventory_plan_docx.py

Requires: python-docx
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent
OUT_PATH = ROOT / "Universal_Product_Inventory_Management_Plan.docx"


def _p(doc: Document, text: str, *, bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic


def _bullets(doc: Document, items: list[str]) -> None:
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def _numbered(doc: Document, items: list[str]) -> None:
    for it in items:
        doc.add_paragraph(it, style="List Number")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i in range(len(headers)):
            cells[i].text = row[i] if i < len(row) else ""
    doc.add_paragraph()


def build_document() -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = doc.add_heading(
        "Universal Product & Inventory Management System",
        0,
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(
        "Planning & Implementation Strategy Document\n"
        f"Generated {datetime.now():%Y-%m-%d %H:%M}\n"
        "Database: Microsoft SQL Server · Extend existing Billing catalog\n"
        "Source: Enterprise AI specification (metadata-driven catalog + ledger-based inventory)"
    )
    run.italic = True
    run.font.size = Pt(10)
    doc.add_paragraph()

    doc.add_heading("1. Purpose of this document", level=1)
    _p(
        doc,
        "This Word document records how we interpret the universal product/inventory AI specification, "
        "how it applies to the Billing workspace that already implements catalog and stock fields on SQL Server, "
        "and exactly what we will do next: extend—not replace—the current system with dynamic metadata, variants, "
        "and a transaction ledger. No PostgreSQL; all persistence is Microsoft SQL Server using the same migration "
        "style as existing dbo scripts.",
    )
    doc.add_page_break()

    doc.add_heading("2. Technology mandate: Microsoft SQL Server", level=1)
    _p(
        doc,
        "All new tables, indexes, and JSON storage use SQL Server features (NVARCHAR(MAX) with JSON validation via "
        "application layer or CHECK constraints, OPENJSON / JSON_VALUE / JSON_QUERY where appropriate, filtered "
        "indexes, columnstore only if justified by reporting volume). Migrations live alongside existing scripts under "
        "db/migrations with idempotent IF OBJECT_ID patterns.",
    )
    _bullets(
        doc,
        [
            "Runtime: ASP.NET Core Web API (existing BillingApp.Api).",
            "Data access: Dapper + IDbConnectionFactory / SqlServerConnectionFactory (already wired).",
            "Auth: JWT + workplace (TenantId, OrganizationId) scoping—reuse on all new tables.",
            "Frontend: React (web/) — extend Products area with dynamic form routes; reports gain search only.",
        ],
    )

    doc.add_heading("3. What already exists (baseline we extend)", level=1)
    _p(
        doc,
        "The repository already implements product and inventory-related catalog behavior. Universal features must "
        "build on this baseline so invoices, purchases, and the current Products UI keep working during rollout.",
    )
    add_table(
        doc,
        ["Area", "Current state (representative)"],
        [
            ["Database", "dbo.Products with TenantId, OrganizationId; SKU, Name, UnitPrice, StockQty, CategoryId, BrandId, UnitId, tax/GST fields, images, etc. (see db/migrations)"],
            ["Backend", "ProductRepository (Dapper), ProductsController under api/v1/products, category/brand/unit repos"],
            ["Stock today", "StockQty column on product — scalar, not variant-level, not ledgered"],
            ["Frontend", "ProductsPage — browse/edit catalog; not metadata-driven forms yet"],
        ],
    )
    doc.add_heading("3.1 Extension principles", level=2)
    _bullets(
        doc,
        [
            "Additive T-SQL only for Phase 1–2: new tables + nullable columns on Products (e.g. HasVariants, DefaultWarehouseId) — never drop StockQty until ledger backfill is verified.",
            "Dual-write or migration window: optionally mirror movements into InventoryTransaction while continuing to update StockQty for simple SKUs until cutover.",
            "APIs remain backward compatible; new endpoints for dynamic-form and variants sit beside existing CRUD.",
        ],
    )
    doc.add_page_break()

    doc.add_heading("4. Analysis of the AI prompt", level=1)
    doc.add_heading("4.1 Core intent", level=2)
    _p(
        doc,
        "Universal product model extended through metadata (dynamic fields + variant axes). Inventory changes only "
        "through an append-only ledger; multi-warehouse, batch/expiry, scale.",
    )
    doc.add_heading("4.2 Non-negotiable architecture rules", level=2)
    _bullets(
        doc,
        [
            "No category-fixed static forms; no RAM/Color/Organic columns on dbo.Products.",
            "Dynamic fields stored as definitions + values (EAV or JSON rows); SQL Server schema stays universal.",
            "Forms driven by GET …/dynamic-form JSON — React renders dynamically.",
            "Sellable stock authoritative at variant (+ warehouse + optional batch) once variants ship.",
        ],
    )
    doc.add_heading("4.3 Mapping prompt → SQL Server", level=2)
    _bullets(
        doc,
        [
            "“JSONB” in prompt → NVARCHAR(MAX) JSON documents or normalized key/value tables with indexed persisted facets for hot filters.",
            "“GIN index” → filtered nonclustered indexes on computed columns / indexed views / full-text if needed.",
            "Redis (prompt): optional cache for session/catalog reads — not required for MVP ledger.",
        ],
    )
    doc.add_page_break()

    doc.add_heading("5. What we will do next (phased checklist)", level=1)
    _p(
        doc,
        "Ordered actions. Each phase extends the Billing codebase and SQL Server schema. Sample seeds prove Samsung, "
        "banana, shirt, grocery, furniture scenarios without new fixed columns.",
    )

    phases: list[tuple[str, list[str]]] = [
        (
            "Phase A — SQL Server schema (additive)",
            [
                "Author T-SQL migration scripts: ProductVariant, ProductCustomFieldDefinition, ProductCustomFieldValue, ProductVariantAttribute, Warehouse, InventoryBalance, InventoryTransaction, Batch, UnitConversion (extend existing Unit master if present), ProductDocument.",
                "Alter dbo.Products: add HasVariants, optional Barcode, PrimaryVariantId nullable FK, JSON metadata bag column only if needed for migration — avoid duplicating vertical facts as columns.",
                "Indexes: SKU/barcode per org, IX on InventoryTransaction (VariantId, WarehouseId, CreatedUtc), filtered indexes for active variants.",
                "ER diagram: Product 1—N Variant; ledger lines reference VariantId (+ WarehouseId, BatchId optional).",
            ],
        ),
        (
            "Phase B — Dynamic field engine (backend)",
            [
                "C# models + repository for field definitions (JSON schema: label, type, validation rules, flags, sectionName, displayOrder).",
                "Persist values per ProductId / VariantId; validate server-side from stored JSON rules.",
                "Integrate with existing ProductRepository.Create/Update paths without breaking current DTOs.",
            ],
        ),
        (
            "Phase C — Dynamic form API",
            [
                "GET /api/v1/products/{id}/dynamic-form — merge system fields (from existing product row) + dynamic sections.",
                "GET /api/v1/products/dynamic-form-template for create-empty wizard.",
                "Map field types to renderer contracts for React.",
            ],
        ),
        (
            "Phase D — Variant engine",
            [
                "Proposal API for Cartesian axes; persist confirmed ProductVariant + ProductVariantAttribute rows.",
                "Price/barcode/weight per variant; images per variant optional.",
                "Business rule: when HasVariants=1, invoices/picks resolve line items to VariantId.",
            ],
        ),
        (
            "Phase E — Inventory ledger",
            [
                "InventoryTransaction types: PURCHASE, SALE, RETURN, DAMAGE, TRANSFER, ADJUSTMENT.",
                "InventoryBalance maintained in same transaction as ledger insert (SERIALIZABLE or row versioning as designed).",
                "Wire existing purchase post / invoice flows to emit ledger lines; define plan to reconcile StockQty vs sum(ledger) before retiring StockQty.",
            ],
        ),
        (
            "Phase F — Units & conversions",
            [
                "Leverage existing Unit table; add UnitConversion with cycle detection.",
                "Optional: normalize ledger quantities to base UOM.",
            ],
        ),
        (
            "Phase G — Modular monolith boundaries",
            [
                "Catalog handlers: products, variants, images, documents, custom fields.",
                "Inventory handlers: warehouses, balances, transactions.",
                "Policies: catalog.manage, inventory.adjust, inventory.transfer.",
            ],
        ),
        (
            "Phase H — Frontend",
            [
                "New routes: catalog wizard (/catalog/products/new, /catalog/products/:id/edit) using DynamicFormRenderer.",
                "Inventory ops pages under /inventory/*.",
                "Existing reports: add search only (debounced, URL sync); no visual redesign.",
            ],
        ),
        (
            "Phase I — Search / facets",
            [
                "Persist searchable flags; SQL Server persisted computed columns or facet table for RAM/Storage-style filters.",
                "POST /catalog/search with DSL + paging.",
            ],
        ),
        (
            "Phase J — Samples & performance",
            [
                "Seed scripts for prompt examples (metadata-only attributes).",
                "Bulk insert tests; consider partitioning InventoryTransaction by date if volume warrants.",
            ],
        ),
    ]

    for title_p, actions in phases:
        doc.add_heading(title_p, level=2)
        _numbered(doc, actions)

    doc.add_page_break()

    doc.add_heading("6. What we implement & how we manage it", level=1)
    doc.add_heading("6.1 Product core + metadata", level=2)
    _p(
        doc,
        "dbo.Products retains universal columns already in production; new attributes use definition/value tables or "
        "JSON blobs in SQL Server. Governance: lock definitions after publish; audit columns on sensitive changes.",
    )
    doc.add_heading("6.2 Variants & inventory", level=2)
    _p(
        doc,
        "Variants become first-class for stock. Legacy single-SKU products map to one implicit variant during migration "
        "so ledger code paths stay unified.",
    )
    add_table(
        doc,
        ["Bucket", "Meaning"],
        [
            ["Available", "Sellable on hand"],
            ["Reserved", "Allocated to orders"],
            ["Damaged", "Write-off"],
            ["Returned", "Returns queue"],
            ["Incoming", "Inbound expected"],
        ],
    )
    doc.add_heading("6.3 Reporting", level=2)
    _p(
        doc,
        "Report pages unchanged visually; add search/filter parameters backed by same query patterns as catalog search.",
    )
    doc.add_page_break()

    doc.add_heading("7. Database summary (SQL Server)", level=1)
    add_table(
        doc,
        ["Table", "Role"],
        [
            ["dbo.Products (extended)", "Universal master; existing columns preserved"],
            ["ProductVariant", "SKU-level sellable row"],
            ["ProductCustomFieldDefinition / Value", "Dynamic schema & answers"],
            ["ProductVariantAttribute", "Axis key/value"],
            ["Warehouse, Batch", "Multi-site & expiry"],
            ["InventoryBalance", "Snapshot"],
            ["InventoryTransaction", "Append-only ledger"],
            ["UnitConversion", "UOM graph"],
        ],
    )
    _p(doc, "Forbidden on Products: RAM, Color, Size, Organic as dedicated columns — metadata only.", italic=True)

    doc.add_page_break()

    doc.add_heading("8. API surface (representative)", level=1)
    add_table(
        doc,
        ["Area", "Examples"],
        [
            ["Catalog", "Existing /api/v1/products/* + GET dynamic-form"],
            ["Variants", "POST variants/generate-proposal; POST variants"],
            ["Inventory", "POST /inventory/transactions"],
            ["Search", "POST /catalog/search"],
        ],
    )

    doc.add_page_break()

    doc.add_heading("9. UI routing", level=1)
    add_table(
        doc,
        ["Route", "Purpose"],
        [
            ["/products (existing)", "Transition to list hub; link to new wizard"],
            ["/catalog/products/new", "Universal create wizard"],
            ["/catalog/products/:id/edit", "Edit + dynamic sections"],
            ["/inventory/*", "Ledger / warehouses"],
            ["/reports/*", "+ search only"],
        ],
    )

    doc.add_heading("10. Risks & mitigations", level=1)
    _bullets(
        doc,
        [
            "StockQty vs ledger drift → reconciliation job + feature flag before dropping StockQty.",
            "JSON query cost → persisted facets / indexed computed columns on SQL Server.",
            "Invoice line migration → resolve ProductId to VariantId at posting time.",
        ],
    )

    doc.add_heading("11. Deliverables", level=1)
    _bullets(
        doc,
        [
            "db/migrations/*.sql for SQL Server",
            "Dapper repositories + API controllers",
            "React DynamicFormRenderer + wizard routes",
            "Integration tests for ledger invariants",
        ],
    )

    doc.add_heading("12. Sign-off", level=1)
    _numbered(
        doc,
        [
            "Confirm SQL Server-only persistence.",
            "Confirm extend-existing-products strategy (no big-bang rewrite).",
            "Approve variant + ledger rollout plan.",
        ],
    )

    return doc


def main() -> None:
    doc = build_document()
    doc.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
