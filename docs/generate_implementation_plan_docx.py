"""
Build Billing_App_Implementation_Plan_V1.docx from docs/IMPLEMENTATION_PLAN_V1.md
Run: python docs/generate_implementation_plan_docx.py
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "IMPLEMENTATION_PLAN_V1.md"


def _strip_inline(md: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", md)


def _add_paragraph(doc: Document, text: str, style: str | None = None):
    text = _strip_inline(text.strip())
    if not text:
        return
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.add_run(text)


def _parse_table(lines: list[str]) -> tuple[list[str], list[list[str]]] | None:
    if len(lines) < 2:
        return None
    def split_row(s: str) -> list[str]:
        parts = [c.strip() for c in s.strip().strip("|").split("|")]
        return parts
    rows = [split_row(l) for l in lines if l.strip().startswith("|")]
    if len(rows) < 2:
        return None

    def is_sep_row(cells: list[str]) -> bool:
        return bool(cells) and all(re.match(r"^:?-+:?$", c.strip()) for c in cells if c.strip())

    if is_sep_row(rows[1]):
        rows = [rows[0]] + rows[2:]
    headers = rows[0]
    body = rows[1:]
    return headers, body


def markdown_file_to_docx(md_file: Path, out_file: Path):
    lines = md_file.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Billing Application - Implementation Plan", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        f"Version 1.0 - Generated {datetime.now():%Y-%m-%d}\nSource markdown: IMPLEMENTATION_PLAN_V1.md"
    )
    r.italic = True
    r.font.size = Pt(10)
    doc.add_paragraph()

    i = 0
    in_code = False
    code_buf: list[str] = []

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_buf = []
            else:
                in_code = False
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if line.strip() == "---":
            i += 1
            continue

        if line.startswith("# "):
            doc.add_heading(_strip_inline(line[2:]), level=0)
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(_strip_inline(line[3:]), level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(_strip_inline(line[4:]), level=2)
            i += 1
            continue
        if line.startswith("#### "):
            doc.add_heading(_strip_inline(line[5:]), level=3)
            i += 1
            continue

        if line.strip().startswith("|"):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl_lines.append(lines[i])
                i += 1
            parsed = _parse_table(tbl_lines)
            if parsed:
                headers, body = parsed
                table = doc.add_table(rows=1, cols=len(headers))
                table.style = "Table Grid"
                for ci, h in enumerate(headers):
                    table.rows[0].cells[ci].text = _strip_inline(h)
                for row in body:
                    cells = table.add_row().cells
                    for ci in range(len(headers)):
                        cells[ci].text = _strip_inline(row[ci]) if ci < len(row) else ""
            continue

        if line.strip().startswith("- "):
            _add_paragraph(doc, line.strip()[2:], style="List Bullet")
            i += 1
            continue

        if re.match(r"^\d+\.\s", line.strip()):
            _add_paragraph(doc, re.sub(r"^\d+\.\s*", "", line.strip()), style="List Number")
            i += 1
            continue

        # ASCII diagram block: preserve monospace if indented with spaces
        if line.startswith(" ") and line.strip():
            block = [line.rstrip()]
            i += 1
            while i < len(lines) and (
                lines[i].startswith(" ")
                or lines[i].strip() in {"?", "?", "?", "?", "?", "?", "?", "?", "?", "?"}
                or "?" in lines[i]
                or "?" in lines[i]
            ):
                block.append(lines[i].rstrip())
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(block))
            r.font.name = "Consolas"
            r.font.size = Pt(8)
            continue

        if line.strip():
            _add_paragraph(doc, line)
        i += 1

    fallbacks = [
        out_file,
        ROOT / "Billing_App_Implementation_Plan_V1_build.docx",
        ROOT / "Billing_App_Implementation_Plan_V1_latest.docx",
    ]
    for path in fallbacks:
        try:
            doc.save(path)
            print(f"Wrote: {path}")
            return path
        except PermissionError:
            continue
    stamped = ROOT / f"Billing_App_Implementation_Plan_V1_{datetime.now():%Y%m%d_%H%M%S}.docx"
    doc.save(stamped)
    print(f"Wrote: {stamped}")
    return stamped


def main():
    if not MD_PATH.exists():
        raise SystemExit(f"Missing {MD_PATH}")
    out = ROOT / "Billing_App_Implementation_Plan_V1.docx"
    markdown_file_to_docx(MD_PATH, out)


if __name__ == "__main__":
    main()
